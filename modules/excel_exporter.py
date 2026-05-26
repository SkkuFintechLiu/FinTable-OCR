from typing import Any, Dict, List, Tuple

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


FIELDS: List[Tuple[str, str]] = [
    ("credit_bonds", "信用类债券"),
    ("bank_loans", "银行贷款"),
    ("non_bank_loans", "非银行金融机构贷款"),
    ("other", "其他有息债务"),
    ("total", "合计"),
    ("short_term", "1年以内"),
]


YELLOW = PatternFill(fill_type="solid", fgColor="FFFF00")
ORANGE = PatternFill(fill_type="solid", fgColor="FFA500")
HEADER = PatternFill(fill_type="solid", fgColor="F3F4F6")


def _cell_display(cell: Dict[str, Any]) -> Any:
    v = cell.get("value")
    if v is not None:
        return v
    raw = cell.get("raw")
    if raw is not None:
        return raw
    return None


def build_workbook_bytes(export_payload: Dict[str, Any]) -> bytes:
    years: List[str] = export_payload.get("years") or []
    rows: List[Dict[str, Any]] = export_payload.get("rows") or []

    wb = Workbook()
    ws = wb.active
    ws.title = "宽表"

    header_cells = ["发行人"]
    for y in years:
        for _, label in FIELDS:
            header_cells.append(f"{label}_{y}")
    header_cells.append("数据来源")

    ws.append(header_cells)
    ws.freeze_panes = "A2"

    header_font = Font(bold=True)
    for col in range(1, len(header_cells) + 1):
        c = ws.cell(row=1, column=col)
        c.font = header_font
        c.fill = HEADER
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for r in rows:
        issuer = r.get("issuer") or ""
        years_map = r.get("years") or {}
        sources = r.get("sources") or {}

        line: List[Any] = [issuer]
        flags: List[str] = ["ok"]

        for y in years:
            yd = years_map.get(y) or {}
            for f, _ in FIELDS:
                cell = yd.get(f) or {"value": None, "raw": None, "flag": "missing"}
                line.append(_cell_display(cell))
                flags.append(cell.get("flag") or "missing")

        parts = []
        for y in years:
            src = sources.get(y) or {}
            st = (src.get("source_type") or "").strip()
            sf = (src.get("source_file") or "").strip()
            notes = (src.get("notes") or "").strip()
            if not st and not sf and not notes:
                continue
            seg = " ".join([x for x in [st, sf] if x])
            if notes:
                seg = f"{seg}（{notes}）" if seg else notes
            if seg:
                parts.append(seg)
        line.append("；".join(dict.fromkeys(parts)))
        flags.append("ok")

        ws.append(line)
        row_idx = ws.max_row
        for col_idx in range(2, len(line)):
            flag = flags[col_idx - 1]
            c = ws.cell(row=row_idx, column=col_idx)
            c.alignment = Alignment(horizontal="center", vertical="center")
            if flag == "missing":
                c.fill = YELLOW
            elif flag in {"invalid", "unit_unknown"}:
                c.fill = ORANGE

    ws.column_dimensions["A"].width = 38
    for i in range(2, len(header_cells) + 1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = 16

    for row in ws.iter_rows(min_row=2, min_col=2, max_col=1 + len(years) * len(FIELDS)):
        for c in row:
            if isinstance(c.value, (int, float)):
                c.number_format = "0.00"

    _build_long_sheet(wb, years, rows)
    _build_error_sheet(wb, export_payload)
    _build_parse_log_sheet(wb, export_payload)

    from io import BytesIO

    bio = BytesIO()
    wb.save(bio)
    return bio.getvalue()


def _build_long_sheet(wb: Workbook, years: List[str], rows: List[Dict[str, Any]]) -> None:
    ws = wb.create_sheet("长表")
    header = ["issuer", "year", "debt_type", "value", "flag", "confidence", "page", "bbox", "parser", "raw_text"]
    ws.append(header)

    header_font = Font(bold=True)
    for col in range(1, len(header) + 1):
        c = ws.cell(row=1, column=col)
        c.font = header_font
        c.fill = HEADER
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for r in rows:
        issuer = r.get("issuer") or ""
        years_map = r.get("years") or {}
        for y in years:
            yd = years_map.get(y) or {}
            trace = (yd.get("__trace__") or {}).get("task") if isinstance(yd.get("__trace__"), dict) else None
            page = trace.get("page") if isinstance(trace, dict) else None
            bbox = trace.get("bbox") if isinstance(trace, dict) else None
            parser = trace.get("parser") if isinstance(trace, dict) else None

            for f, label in FIELDS:
                cell = yd.get(f) or {"value": None, "raw": None, "flag": "missing"}
                v = cell.get("value")
                raw = cell.get("raw")
                flag = cell.get("flag") or "missing"
                conf = cell.get("confidence")
                ws.append([issuer, y, label, v, flag, conf, page, str(bbox) if bbox else "", parser, raw])

    ws.freeze_panes = "A2"
    ws.column_dimensions["A"].width = 38
    for col_letter in ["B", "C", "D", "E", "F", "G", "H", "I", "J"]:
        ws.column_dimensions[col_letter].width = 18
    for row in ws.iter_rows(min_row=2, min_col=4, max_col=4):
        for c in row:
            if isinstance(c.value, (int, float)):
                c.number_format = "0.00"


def _build_error_sheet(wb: Workbook, export_payload: Dict[str, Any]) -> None:
    ws = wb.create_sheet("错误日志")
    header = ["issuer", "year", "field", "flag", "raw_text", "source"]
    ws.append(header)

    header_font = Font(bold=True)
    for col in range(1, len(header) + 1):
        c = ws.cell(row=1, column=col)
        c.font = header_font
        c.fill = HEADER
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    years: List[str] = export_payload.get("years") or []
    rows: List[Dict[str, Any]] = export_payload.get("rows") or []
    for r in rows:
        issuer = r.get("issuer") or ""
        years_map = r.get("years") or {}
        sources = r.get("sources") or {}
        for y in years:
            yd = years_map.get(y) or {}
            src = sources.get(y) or {}
            source = " ".join([x for x in [(src.get("source_type") or "").strip(), (src.get("source_file") or "").strip()] if x])
            notes = (src.get("notes") or "").strip()
            if notes:
                source = f"{source}（{notes}）" if source else notes
            for f, label in FIELDS:
                cell = yd.get(f) or {"value": None, "raw": None, "flag": "missing"}
                flag = cell.get("flag") or "missing"
                if flag == "ok":
                    continue
                ws.append([issuer, y, label, flag, cell.get("raw") or "", source])

    ws.freeze_panes = "A2"
    ws.column_dimensions["A"].width = 38
    for col_letter in ["B", "C", "D", "E", "F"]:
        ws.column_dimensions[col_letter].width = 22


def _build_parse_log_sheet(wb: Workbook, export_payload: Dict[str, Any]) -> None:
    ws = wb.create_sheet("解析日志")
    header = ["issuer", "year", "page", "parser", "bbox", "title_snippet"]
    ws.append(header)

    header_font = Font(bold=True)
    for col in range(1, len(header) + 1):
        c = ws.cell(row=1, column=col)
        c.font = header_font
        c.fill = HEADER
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    years: List[str] = export_payload.get("years") or []
    rows: List[Dict[str, Any]] = export_payload.get("rows") or []
    for r in rows:
        issuer = r.get("issuer") or ""
        years_map = r.get("years") or {}
        for y in years:
            yd = years_map.get(y) or {}
            trace = yd.get("__trace__") if isinstance(yd.get("__trace__"), dict) else None
            task_trace = trace.get("task") if isinstance(trace, dict) else None
            if not isinstance(task_trace, dict):
                continue
            ws.append(
                [
                    issuer,
                    y,
                    task_trace.get("page") or "",
                    task_trace.get("parser") or "",
                    str(task_trace.get("bbox") or ""),
                    task_trace.get("text") or "",
                ]
            )

    ws.freeze_panes = "A2"
    ws.column_dimensions["A"].width = 38
    for col_letter in ["B", "C", "D", "E", "F"]:
        ws.column_dimensions[col_letter].width = 22
